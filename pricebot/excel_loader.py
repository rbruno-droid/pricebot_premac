from __future__ import annotations
import os
import re
import tempfile
from dataclasses import dataclass
from typing import List, Tuple, Dict, Optional, Any

import pandas as pd
from openpyxl import load_workbook

from .text_utils import normalize_text, normalize_code

STANDARD_COLUMNS = [
    "codigo",
    "nombre_producto",
    "descripcion",
    "precio_unitario",
    "precio_3_5",
    "precio_6_mas",
    "moneda",
    "fuente_archivo",
    "fuente_hoja",
    "fila_origen",
]

HEADER_ALIASES = {
    "codigo": ["codigo", "código", "numero de parte", "número de parte", "parte", "part number", "p/n", "pn", "referencia", "ref"],
    "nombre_producto": ["nombre del producto", "producto", "product name", "nombre producto", "nombre"],
    "descripcion": ["descripcion", "descripción", "description", "detalle", "detalle producto"],
    "precio_unitario": ["precio unitario con descuento (1-2 unidades)", "precio unitario con descuento", "precio unitario", "precio", "unit price", "price"],
    "precio_3_5": ["precio unitario con descuento (3-5 unidades)", "precio 3-5", "3-5 unidades", "precio_3_5"],
    "precio_6_mas": ["precio unitario con descuento (6+ unidades)", "precio 6+", "6+ unidades", "precio_6_mas"],
}

BAD_CODE_VALUES = {"codigo", "código", "numero de parte", "número de parte", "part number", "p/n", "pn", "referencia", "ref"}
HELP_SHEET_MARKERS = ["ayuda", "help", "instrucciones"]

@dataclass
class LoadReport:
    source: str
    sheet: str
    status: str
    rows_loaded: int = 0
    message: str = ""
    header_row: Optional[int] = None
    detected_columns: Optional[Dict[str, str]] = None


def _norm_header(value: Any) -> str:
    return normalize_text(value).replace("_", " ")


def _find_header_row_from_rows(rows, max_scan_rows: int = 80) -> Tuple[Optional[int], Dict[str, int], Dict[str, str]]:
    """Busca una fila de encabezados desde una matriz de valores. Devuelve fila 1-based y columnas 0-based."""
    best = (None, {}, {}, 0)
    for r_idx, row in enumerate(rows[:max_scan_rows], start=1):
        normalized = [_norm_header(v) for v in row]
        mapping: Dict[str, int] = {}
        labels: Dict[str, str] = {}
        for std, aliases in HEADER_ALIASES.items():
            for c_idx, hv in enumerate(normalized):
                if not hv:
                    continue
                for a in aliases:
                    na = _norm_header(a)
                    if hv == na or na in hv:
                        mapping[std] = c_idx
                        labels[std] = str(row[c_idx])
                        break
                if std in mapping:
                    break
        score = 0
        if "codigo" in mapping: score += 3
        if "nombre_producto" in mapping: score += 2
        if "descripcion" in mapping: score += 1
        if "precio_unitario" in mapping: score += 3
        if score > best[3]:
            best = (r_idx, mapping, labels, score)
    if best[3] >= 8 and "codigo" in best[1] and "nombre_producto" in best[1] and "precio_unitario" in best[1]:
        return best[0], best[1], best[2]
    return None, {}, {}

def _to_number(value: Any):
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if pd.isna(value):
            return None
        return float(value)
    s = str(value).strip()
    if not s or normalize_text(s) in ["n/a", "na", "none", "null", "-", "obsolete", "obsoleto"]:
        return None
    # Remueve moneda y separadores. Soporta 1.234,56 y 1,234.56 de forma básica.
    s2 = re.sub(r"[^0-9,\.\-]", "", s)
    if not s2:
        return None
    if "," in s2 and "." in s2:
        if s2.rfind(",") > s2.rfind("."):
            s2 = s2.replace(".", "").replace(",", ".")
        else:
            s2 = s2.replace(",", "")
    elif "," in s2 and "." not in s2:
        s2 = s2.replace(",", ".")
    try:
        return float(s2)
    except Exception:
        return None


def _is_probably_header_or_section(code, name, desc, price):
    code_s = normalize_text(code)
    name_s = normalize_text(name)
    desc_s = normalize_text(desc)
    if not code_s:
        return True
    if code_s in BAD_CODE_VALUES:
        return True
    # Fila de categoría/sección: trae texto largo en columna código, sin producto/descripción/precio.
    if price is None and not name_s and not desc_s and len(code_s.split()) >= 3:
        return True
    return False


def load_excel_file(path: str, source_name: Optional[str] = None) -> Tuple[pd.DataFrame, List[LoadReport]]:
    source_name = source_name or os.path.basename(path)
    reports: List[LoadReport] = []
    records: List[dict] = []
    wb = load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        sheet_name = ws.title
        if any(m in normalize_text(sheet_name) for m in HELP_SHEET_MARKERS):
            reports.append(LoadReport(source_name, sheet_name, "omitida", 0, "Hoja de ayuda/instrucciones"))
            continue
        # En modo read_only, usar iter_rows es mucho más rápido que ws.cell(r,c).
        rows = list(ws.iter_rows(values_only=True))
        header_row, mapping, labels = _find_header_row_from_rows(rows)
        if not header_row:
            reports.append(LoadReport(source_name, sheet_name, "sin_encabezado", 0, "No encontré encabezados estándar"))
            continue
        loaded = 0
        for r_idx, row in enumerate(rows[header_row:], start=header_row + 1):
            def val(std):
                c = mapping.get(std)
                return row[c] if c is not None and c < len(row) else None
            code = val("codigo")
            name = val("nombre_producto")
            desc = val("descripcion")
            p1 = _to_number(val("precio_unitario"))
            p35 = _to_number(val("precio_3_5"))
            p6 = _to_number(val("precio_6_mas"))
            if _is_probably_header_or_section(code, name, desc, p1):
                continue
            rec = {
                "codigo": str(code).strip() if code is not None else "",
                "nombre_producto": str(name).strip() if name is not None else "",
                "descripcion": str(desc).strip() if desc is not None else "",
                "precio_unitario": p1,
                "precio_3_5": p35,
                "precio_6_mas": p6,
                "moneda": "USD",
                "fuente_archivo": source_name,
                "fuente_hoja": sheet_name,
                "fila_origen": r_idx,
            }
            rec["codigo_normalizado"] = normalize_code(rec["codigo"])
            rec["texto_busqueda"] = " ".join([
                rec["codigo"], rec["codigo_normalizado"], rec["nombre_producto"], rec["descripcion"], source_name, sheet_name
            ]).strip()
            records.append(rec)
            loaded += 1
        reports.append(LoadReport(source_name, sheet_name, "ok", loaded, "Cargada", header_row, labels))
    df = pd.DataFrame(records)
    if df.empty:
        return pd.DataFrame(columns=STANDARD_COLUMNS + ["codigo_normalizado", "texto_busqueda"]), reports
    return df, reports

def choose_price(row: pd.Series, qty: int):
    if qty >= 6 and pd.notna(row.get("precio_6_mas")):
        return float(row.get("precio_6_mas")), "6+ unidades"
    if 3 <= qty <= 5 and pd.notna(row.get("precio_3_5")):
        return float(row.get("precio_3_5")), "3-5 unidades"
    if pd.notna(row.get("precio_unitario")):
        return float(row.get("precio_unitario")), "1-2 unidades / precio unitario"
    return None, "sin precio"
